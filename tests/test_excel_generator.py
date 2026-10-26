"""Unit tests for Excel generation, binary signatures, openpyxl styles, and SUM formulas."""

import io

import openpyxl

from app.excel.generator import excel_generator
from app.schemas import TransactionItem


def test_excel_generator_binary_signature(sample_transactions_list: list[TransactionItem]) -> None:
    """Ensure the generated Excel stream starts with the official PK zip magic header."""
    stream = excel_generator.generate_reconciliation_workbook(
        sample_transactions_list, report_title="Test Conciliacion"
    )
    content = stream.getvalue()

    # XLSX is a ZIP package starting with b'PK\x03\x04'
    assert content.startswith(b"PK\x03\x04")
    assert len(content) > 1000


def test_excel_generator_structure_and_formulas(
    sample_transactions_list: list[TransactionItem],
) -> None:
    """Verify sheets, headers, data rows, and dynamic SUM formulas."""
    stream = excel_generator.generate_reconciliation_workbook(
        sample_transactions_list, report_title="Auditoria Q2"
    )
    wb = openpyxl.load_workbook(io.BytesIO(stream.getvalue()))

    # Verify sheet names
    assert "Transacciones" in wb.sheetnames
    assert "Resumen Ejecutivo" in wb.sheetnames

    ws_tx = wb["Transacciones"]

    # Header check
    assert ws_tx.cell(row=1, column=1).value == "#"
    assert ws_tx.cell(row=1, column=2).value == "Fecha"
    assert ws_tx.cell(row=1, column=3).value == "Concepto / Descripción"
    assert ws_tx.cell(row=1, column=6).value == "Cargo (Egreso)"
    assert ws_tx.cell(row=1, column=7).value == "Abono (Ingreso)"

    # Font family check
    assert ws_tx.cell(row=1, column=1).font.name == "Segoe UI"
    assert ws_tx.cell(row=1, column=1).font.bold is True

    # Data rows check (4 transactions -> rows 2 to 5)
    num_tx = len(sample_transactions_list)
    assert ws_tx.cell(row=2, column=2).value == sample_transactions_list[0].fecha

    # Check totals row (row 6)
    total_row = num_tx + 2
    assert ws_tx.cell(row=total_row, column=1).value == "TOTALES"

    formula_cargo = ws_tx.cell(row=total_row, column=6).value
    formula_abono = ws_tx.cell(row=total_row, column=7).value

    assert formula_cargo == f"=SUM(F2:F{num_tx + 1})"
    assert formula_abono == f"=SUM(G2:G{num_tx + 1})"

    # Summary sheet verification
    ws_summary = wb["Resumen Ejecutivo"]
    assert "Auditoria Q2" in str(ws_summary.cell(row=2, column=2).value)
    assert ws_summary.cell(row=5, column=2).value == "Métrica Financiera"


def test_excel_generator_without_summary(sample_transactions_list: list[TransactionItem]) -> None:
    """Verify that include_summary=False produces only the transactions sheet."""
    stream = excel_generator.generate_reconciliation_workbook(
        sample_transactions_list, include_summary=False
    )
    wb = openpyxl.load_workbook(io.BytesIO(stream.getvalue()))

    assert wb.sheetnames == ["Transacciones"]
